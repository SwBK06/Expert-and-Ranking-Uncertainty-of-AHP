from pathlib import Path
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
import matplotlib.colors as mcolors
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from sklearn import base
from sklearn.decomposition import PCA
from sklearn.preprocessing import StandardScaler
from scipy.stats import rankdata, spearmanr, kendalltau, qmc
FILE_PATH = Path(
    r"C:\Users\Swarna\Desktop\PhD\Analysis\Part 1\Python"
    r"\Sensitivity1y - old (for discover hazard).xlsx"
)
OUTPUT_DIR = FILE_PATH.parent
SHEET_NAME, TARGET_COL = (0, 'target_col')
OUTPUT_FILE = OUTPUT_DIR / 'SAW_AHP_PCA_sensitivity.xlsx'
OAT_DELTAS = (0.1, 0.2, 0.3)
N_LHS, LHS_RANGE, LHS_SEED = (2000, 0.2, 42)
TOP_PCT, TOP_K = (0.2, 36)
RUN_STOCHASTIC_DATA_SA = True
DATA_ERROR_LEVELS = (0.05, 0.1, 0.15, 0.2)
N_DATA_SIM = 1000
DATA_SEED = 42
ADD_CMAP, SUB_CMAP = ('Oranges', 'Purples')
CRITERIA = ['benefit', 'cost', 'cost', 'cost', 'cost', 'benefit', 'benefit', 'benefit', 'cost', 'cost', 'cost', 'cost', 'cost', 'cost', 'benefit', 'cost', 'benefit']
AHP_WEIGHTS = np.array([0.0292,0.0586,0.0382,0.0425,0.0512,0.0303,0.0301,0.0207,0.0149,0.0401,0.0376,0.0200,0.0192,0.0233,0.0242,0.0199,0.5000], dtype=float)
MAX_VALS = {'C1': 1.0, 'C2': 0.4168, 'C3': 6.6916, 'C4': 0.2563, 'C5': 1.0, 'C6': 41888.9322, 'C7': 4.39, 'C8': 2.02, 'C9': 100.0, 'C10': 100.0, 'C11': 100.0, 'C12': 100.0, 'C13': 100.0, 'C14': 1.0, 'C15': 1.0, 'C16': 100.0, 'C17': 60.0}
MIN_VALS = {'C1': 0.0, 'C2': 0.0, 'C3': 0.0, 'C4': 0.0, 'C5': 0.0, 'C6': 0.0, 'C7': 1.0, 'C8': 0.09, 'C9': 0.0, 'C10': 0.0, 'C11': 0.0, 'C12': 0.0, 'C13': 0.0, 'C14': 0.0, 'C15': 0.0, 'C16': 0.0, 'C17': 0.0}
CRITERION_NAMES = {'C1': 'Settlement', 'C2': 'Health Facility', 'C3': 'School', 'C4': 'Shelter', 'C5': 'Polder Coverage', 'C6': 'Population Density', 'C7': 'Household Size', 'C8': 'Sex Ratio', 'C9': 'Literacy Rate', 'C10': 'Access to Drinking Water Source', 'C11': 'Access to Sanitation Facility', 'C12': 'Access to Power Supply', 'C13': 'Access to Information', 'C14': 'Employment Rate', 'C15': 'Agricultural Labor Share', 'C16': 'Access to Financial Institution', 'C17': 'Flood Extent'}

def load_excel_data(file_path, sheet_name=0, header=0):
    return pd.read_excel(file_path, sheet_name=sheet_name, header=header, engine='openpyxl')

def saw(df, weights, criteria, target_column, norm_method='minmax'):
    eps = 1e-08
    cols = [c for c in df.columns if c != target_column]
    matrix = df[cols].to_numpy(dtype=float)
    weights = np.asarray(weights, dtype=float)
    if len(criteria) != matrix.shape[1] or len(weights) != matrix.shape[1]:
        raise ValueError('Criteria/weights do not match decision columns.')
    norm = np.zeros_like(matrix)
    if norm_method.lower() == 'minmax':
        for j, direction in enumerate(criteria):
            col = matrix[:, j]
            cmin, cmax = (np.nanmin(col), np.nanmax(col))
            denom = cmax - cmin + eps
            if direction.lower() == 'benefit':
                norm[:, j] = (col - cmin) / denom
            elif direction.lower() == 'cost':
                norm[:, j] = (cmax - col) / denom
            else:
                raise ValueError("Criteria must be 'benefit' or 'cost'.")
    elif norm_method.lower() == 'zscore':
        means = np.nanmean(matrix, axis=0)
        stds = np.nanstd(matrix, axis=0) + eps
        norm = (matrix - means) / stds
        for j, direction in enumerate(criteria):
            if direction.lower() == 'cost':
                norm[:, j] *= -1
            elif direction.lower() != 'benefit':
                raise ValueError("Criteria must be 'benefit' or 'cost'.")
    else:
        raise ValueError("norm_method must be 'minmax' or 'zscore'.")
    scores = norm @ weights
    ranks = rankdata(-scores, method='ordinal')
    return (scores, ranks)

def pca_weights(df, target_column):
    x = df.drop(columns=[target_column], errors='ignore')
    x_scaled = StandardScaler().fit_transform(x)
    pca = PCA(n_components=1)
    pca.fit(x_scaled)
    loadings = pca.components_[0]
    abs_loadings = np.abs(loadings)
    weights = abs_loadings / abs_loadings.sum()
    table = pd.DataFrame({'Criterion': x.columns, 'PC1_Loading': loadings, 'Absolute_PC1_Loading': abs_loadings, 'PCA_Weight': weights})
    return (weights, float(pca.explained_variance_ratio_[0]), table)

def ranking_metrics(base_ranks, new_ranks, top_pct=0.2, top_k=20):
    base_ranks, new_ranks = (np.asarray(base_ranks), np.asarray(new_ranks))
    n, k = (len(base_ranks), min(top_k, len(base_ranks)))
    n_top = max(1, int(np.ceil(n * top_pct)))
    base_top = set(np.argsort(base_ranks)[:n_top])
    new_top = set(np.argsort(new_ranks)[:n_top])
    base_k = set(np.argsort(base_ranks)[:k])
    new_k = set(np.argsort(new_ranks)[:k])
    abs_change = np.abs(base_ranks - new_ranks)
    critical_idx = np.array(sorted(base_k), dtype=int)
    union_k = base_k | new_k
    rho, rho_p = spearmanr(base_ranks, new_ranks)
    tau, tau_p = kendalltau(base_ranks, new_ranks)
    return {'Spearman_rho': rho, 'Spearman_p': rho_p, 'Kendall_tau': tau, 'Kendall_p': tau_p, 'Mean_absolute_rank_change': abs_change.mean(), 'Maximum_rank_change': abs_change.max(), 'Total_rank_change': abs_change.sum(), 'Top20_overlap_fraction': len(base_top & new_top) / n_top, 'TopK_Jaccard': len(base_k & new_k) / len(union_k) if union_k else 1.0, 'Critical_mean_abs_rank_change': abs_change[critical_idx].mean(), 'Critical_max_abs_rank_change': abs_change[critical_idx].max()}

def oat_weight_sensitivity(decision_df, target_col, weights, criteria, deltas=OAT_DELTAS, top_pct=TOP_PCT, top_k=TOP_K):
    weights = np.asarray(weights, dtype=float)
    _, base_ranks = saw(decision_df, weights, criteria, target_col)
    cols = [c for c in decision_df.columns if c != target_col]
    add_table, sub_table, rows = ({}, {}, [])
    for delta in deltas:
        add_row, sub_row = ({}, {})
        for j, col in enumerate(cols):
            for sign, direction in ((1, 'Add'), (-1, 'Subtract')):
                w = weights.copy()
                w[j] *= 1 + sign * delta
                w /= w.sum()
                _, ranks = saw(decision_df, w, criteria, target_col)
                metrics = ranking_metrics(base_ranks, ranks, top_pct, top_k)
                metrics.update({'Criterion': col, 'Direction': direction, 'Delta': delta, 'Perturbed_weight': w[j]})
                rows.append(metrics)
                if sign == 1:
                    add_row[col] = metrics['Total_rank_change']
                else:
                    sub_row[col] = metrics['Total_rank_change']
        add_table[delta], sub_table[delta] = (add_row, sub_row)
    return (pd.DataFrame(add_table).T, pd.DataFrame(sub_table).T, pd.DataFrame(rows))

def stochastic_data_sensitivity(decision_df, target_col, weights, criteria, labels, rel_errors=DATA_ERROR_LEVELS, n_sim=1000, seed=42, min_values=None, max_values=None):
    rng = np.random.default_rng(seed)
    _, base_ranks = saw(decision_df, weights, criteria, target_col)
    cols = [c for c in decision_df.columns if c != target_col]
    summary_rows = []
    simulation_rows = []
    for rel_error in rel_errors:
        rank_changes = []
        for sim in range(n_sim):
            perturbed = decision_df.copy()
            for col in cols:
                x = perturbed[col].to_numpy(float)
                noise = rng.uniform(-rel_error, rel_error, size=len(x))
                x_new = x * (1 + noise)
                if min_values and col in min_values:
                    x_new = np.maximum(x_new, min_values[col])
                if max_values and col in max_values:
                    x_new = np.minimum(x_new, max_values[col])
                perturbed[col] = x_new
            _, ranks = saw(perturbed, weights, criteria, target_col)
            abs_change = np.abs(base_ranks - ranks)
            rank_changes.append(abs_change)
            simulation_rows.append({'Perturbation': rel_error, 'Simulation': sim + 1, 'Mean_Abs_Rank_Change': abs_change.mean(), 'Median_Abs_Rank_Change': np.median(abs_change), 'P95_Abs_Rank_Change': np.percentile(abs_change, 95), 'Maximum_Rank_Change': abs_change.max()})
        rank_changes = np.vstack(rank_changes)
        for i, label in enumerate(labels):
            summary_rows.append({'Perturbation': rel_error, 'Alternative': label, 'Baseline_Rank': int(base_ranks[i]), 'Mean_Abs_Rank_Change': rank_changes[:, i].mean(), 'Median_Abs_Rank_Change': np.median(rank_changes[:, i]), 'P95_Abs_Rank_Change': np.percentile(rank_changes[:, i], 95), 'Max_Abs_Rank_Change': rank_changes[:, i].max()})
    return (pd.DataFrame(summary_rows), pd.DataFrame(simulation_rows))

def lhs_weight_uncertainty(decision_df, target_col, weights, criteria, labels, n_sim=N_LHS, rel_range=LHS_RANGE, seed=LHS_SEED, top_pct=TOP_PCT, top_k=TOP_K, preserve_components=False):
    weights = np.asarray(weights, dtype=float)
    labels = np.asarray(labels)
    base_scores, base_ranks = saw(decision_df, weights, criteria, target_col)
    sampler = qmc.LatinHypercube(d=len(weights), seed=seed)
    u = sampler.random(n_sim)
    multipliers = 1 - rel_range + 2 * rel_range * u
    sampled_weights = weights * multipliers
    if preserve_components:
        sampled_weights[:, :6] *= 0.25 / sampled_weights[:, :6].sum(axis=1, keepdims=True)
        sampled_weights[:, 6:16] *= 0.25 / sampled_weights[:, 6:16].sum(axis=1, keepdims=True)
        sampled_weights[:, 16] = 0.5
    else:
        sampled_weights /= sampled_weights.sum(axis=1, keepdims=True)
    metrics_rows = []
    critical_idx = np.argsort(base_ranks)[:min(top_k, len(base_ranks))]
    critical_ranks = np.zeros((n_sim, len(critical_idx)), dtype=float)
    for s, w in enumerate(sampled_weights):
        _, ranks = saw(decision_df, w, criteria, target_col)
        row = ranking_metrics(base_ranks, ranks, top_pct, top_k)
        row['Simulation'] = s + 1
        metrics_rows.append(row)
        critical_ranks[s] = ranks[critical_idx]
    metrics_df = pd.DataFrame(metrics_rows)
    weight_cols = [c for c in decision_df.columns if c != target_col]
    weight_df = pd.DataFrame(sampled_weights, columns=weight_cols)
    weight_df.insert(0, 'Simulation', np.arange(1, n_sim + 1))
    critical_rows = []
    for pos, idx in enumerate(critical_idx):
        sim_ranks = critical_ranks[:, pos]
        abs_change = np.abs(sim_ranks - base_ranks[idx])
        critical_rows.append({'Alternative': labels[idx], 'Baseline_rank': int(base_ranks[idx]), 'Median_simulated_rank': np.median(sim_ranks), 'Mean_absolute_rank_change': abs_change.mean(), 'Median_absolute_rank_change': np.median(abs_change), 'P95_absolute_rank_change': np.percentile(abs_change, 95), 'Maximum_absolute_rank_change': abs_change.max(), 'Probability_remain_in_topK': np.mean(sim_ranks <= min(top_k, len(base_ranks)))})
    summary_df = metrics_df.drop(columns='Simulation').describe(percentiles=[0.05, 0.5, 0.95]).T.reset_index().rename(columns={'index': 'Metric'})
    base_df = pd.DataFrame({'Alternative': labels, 'Baseline_score': base_scores, 'Baseline_rank': base_ranks}).sort_values('Baseline_rank')
    return (metrics_df, weight_df, pd.DataFrame(critical_rows), summary_df, base_df)

def compare_ahp_pca(decision_df, target_col, ahp_weights, pca_weights_, criteria, labels, top_pct=TOP_PCT, top_k=TOP_K):
    ahp_scores, ahp_ranks = saw(decision_df, ahp_weights, criteria, target_col)
    pca_scores, pca_ranks = saw(decision_df, pca_weights_, criteria, target_col)
    metrics = ranking_metrics(ahp_ranks, pca_ranks, top_pct, top_k)
    comparison = pd.DataFrame([{'Comparison': 'AHP-SAW vs PCA-SAW', **metrics}])
    ranks = pd.DataFrame({'Alternative': labels, 'AHP_Score': ahp_scores, 'AHP_Rank': ahp_ranks, 'PCA_Score': pca_scores, 'PCA_Rank': pca_ranks, 'Absolute_rank_change': np.abs(ahp_ranks - pca_ranks)}).sort_values('AHP_Rank')
    weights = pd.DataFrame({'Criterion': [c for c in decision_df.columns if c != target_col], 'AHP_Weight': ahp_weights, 'PCA_Weight': pca_weights_})
    return (comparison, ranks, weights)

def make_component_weights(base_weights, s_share, v_share, h_share):
    """Preserve within-component AHP proportions while changing S/V/H shares."""
    base = np.asarray(base_weights, dtype=float)
    s_internal = base[:6] / base[:6].sum()
    v_internal = base[6:16] / base[6:16].sum()
    w = np.zeros(17, dtype=float)
    w[:6] = s_internal * s_share
    w[6:16] = v_internal * v_share
    w[16] = h_share
    return w / w.sum()

def component_scores(decision_df, target_col, base_weights, criteria):
    """Calculate internally weighted S, V and H component scores."""
    base = np.asarray(base_weights, dtype=float)
    s_weights = base[:6] / base[:6].sum()
    v_weights = base[6:16] / base[6:16].sum()
    s_score, _ = saw(decision_df.iloc[:, :6], s_weights, criteria[:6], target_col)
    v_score, _ = saw(decision_df.iloc[:, 6:16], v_weights, criteria[6:16], target_col)
    h_score, _ = saw(decision_df.iloc[:, 16:17], np.array([1.0]), criteria[16:17], target_col)
    return (s_score, v_score, h_score)

def geometric_component_score(s_score, v_score, h_score, s_weight=0.25, v_weight=0.25, h_weight=0.5, eps=1e-08):
    """Weighted geometric aggregation of S, V and H component scores."""
    score = (s_score + eps) ** s_weight * (v_score + eps) ** v_weight * (h_score + eps) ** h_weight
    ranks = rankdata(-score, method='ordinal')
    return (score, ranks)

def component_sensitivity(decision_df, target_col, labels, base_weights, criteria, top_pct=TOP_PCT, top_k=TOP_K):
    """Run additive structural scenarios plus one geometric aggregation case."""
    base = np.asarray(base_weights, dtype=float)
    base /= base.sum()
    h67 = make_component_weights(base, 0.1665, 0.1665, 0.667)
    h33 = make_component_weights(base, 0.33335, 0.33335, 0.3333)
    h25 = make_component_weights(base, 0.375, 0.375, 0.25)
    s30_v20 = make_component_weights(base, 0.3, 0.2, 0.5)
    s20_v30 = make_component_weights(base, 0.2, 0.3, 0.5)
    base_name = 'S(25%) V(25%) H(50%)'
    scenarios = {base_name: base, 'S(16.65%) V(16.65%) H (66.7%)': h67, 'S(33.35%) V(33.35%) H(33.3%)': h33, 'S(37.5%) V(37.5%) H(25%)': h25, 'S(30%) V(20%) H(50%)': s30_v20, 'S(20%) V(30%) H(50%)': s20_v30}
    results = {}
    for name, w in scenarios.items():
        scores, ranks = saw(decision_df, w, criteria, target_col)
        results[name] = {'scores': scores, 'ranks': ranks, 'weights': w}
    base_ranks = results[base_name]['ranks']
    base_k_idx = set(np.argsort(base_ranks)[:min(top_k, len(base_ranks))])
    s_score, v_score, h_score = component_scores(decision_df, target_col, base, criteria)
    geo_score, geo_rank = geometric_component_score(s_score, v_score, h_score, 0.25, 0.25, 0.5)
    metric_rows, change_rows = ([], [])
    metric_rows.append({'Scenario': 'S(25%) V(25%) H(50%)', 'Spearman_rho': 1.0, 'Spearman_p': 0.0, 'Kendall_tau': 1.0, 'Kendall_p': 0.0, 'Mean_absolute_rank_change': 0.0, 'Maximum_rank_change': 0.0, 'Total_rank_change': 0.0, 'Top20_overlap_fraction': 1.0, 'TopK_Jaccard': 1.0, 'Critical_mean_abs_rank_change': 0.0, 'Critical_max_abs_rank_change': 0.0, 'TopK_Retained_Count': top_k, 'TopK_Dropped_Count': 0, 'TopK_Entered_Count': 0})
    def add_scenario_metrics(name, ranks):
        metrics = ranking_metrics(base_ranks, ranks, top_pct, top_k)
        scenario_k_idx = set(np.argsort(ranks)[:min(top_k, len(ranks))])
        retained = base_k_idx & scenario_k_idx
        dropped = base_k_idx - scenario_k_idx
        entered = scenario_k_idx - base_k_idx
        metrics.update({'Scenario': name, 'TopK_Retained_Count': len(retained), 'TopK_Dropped_Count': len(dropped), 'TopK_Entered_Count': len(entered)})
        metric_rows.append(metrics)
        for idx in sorted(dropped, key=lambda i: base_ranks[i]):
            change_rows.append({'Scenario': name, 'Change': 'Dropped_from_topK', 'Union_Name': labels[idx], 'Baseline_Rank': int(base_ranks[idx]), 'Scenario_Rank': int(ranks[idx])})
        for idx in sorted(entered, key=lambda i: ranks[i]):
            change_rows.append({'Scenario': name, 'Change': 'Entered_topK', 'Union_Name': labels[idx], 'Baseline_Rank': int(base_ranks[idx]), 'Scenario_Rank': int(ranks[idx])})
    for name, result in results.items():
        if name != base_name:
            add_scenario_metrics(name, result['ranks'])
    add_scenario_metrics('Geometric S(25%) V(25%) H(50%)', geo_rank)
    scores_df = pd.DataFrame({'Union_Name': labels})
    for name, result in results.items():
        scores_df[f'{name}_Score'] = result['scores']
        scores_df[f'{name}_Rank'] = result['ranks']
    scores_df['Geometric_Score'] = geo_score
    scores_df['Geometric_Rank'] = geo_rank
    component_rows = []
    for name, result in results.items():
        w = result['weights']
        component_rows.append({'Scenario': name, 'Susceptibility': w[:6].sum(), 'Vulnerability': w[6:16].sum(), 'Hazard': w[16], 'Aggregation': 'Additive SAW'})
    component_rows.append({'Scenario': 'Geometric S(25%) V(25%) H(50%)', 'Susceptibility': 0.25, 'Vulnerability': 0.25, 'Hazard': 0.5, 'Aggregation': 'Weighted geometric'})
    changes_df = pd.DataFrame(change_rows)
    if changes_df.empty:
        changes_df = pd.DataFrame(columns=['Scenario', 'Change', 'Union_Name', 'Baseline_Rank', 'Scenario_Rank'])
    return (scores_df, pd.DataFrame(metric_rows), pd.DataFrame(component_rows), changes_df)

def summarize_stochastic(sim_df):
    return (sim_df.groupby("Perturbation", as_index=False)
            .agg(Mean=("Mean_Abs_Rank_Change", "mean"),
                 Median=("Mean_Abs_Rank_Change", "median"),
                 Q1=("Mean_Abs_Rank_Change", lambda x: np.percentile(x, 25)),
                 Q3=("Mean_Abs_Rank_Change", lambda x: np.percentile(x, 75)),
                 P05=("Mean_Abs_Rank_Change", lambda x: np.percentile(x, 5)),
                 P95=("Mean_Abs_Rank_Change", lambda x: np.percentile(x, 95)),
                 Maximum=("Mean_Abs_Rank_Change", "max"))
            .sort_values("Perturbation"))

def plot_stochastic_data_sensitivity_combined(
    stochastic_ahp_df, stochastic_pca_df,
    save_path="AHP_PCA_Stochastic_Data_Sensitivity.png", figsize=(10, 5)
):
    fig, axes = plt.subplots(1, 2, figsize=figsize, sharey=True)
    datasets = [(stochastic_ahp_df, "(a) AHP"),
                (stochastic_pca_df, "(b) PCA")]
    for ax, (df, title) in zip(axes, datasets):
        summary = (df.groupby("Perturbation", as_index=False)
                   .agg(Mean=("Mean_Abs_Rank_Change", "mean"),
                        Maximum=("Max_Abs_Rank_Change", "max"))
                   .sort_values("Perturbation"))
        x = np.arange(len(summary))
        labels = [f"±{int(v * 100)}%" for v in summary["Perturbation"]]
        width = 0.35
        bars1 = ax.bar(x - width / 2, summary["Mean"], width, label="Mean")
        bars2 = ax.bar(x + width / 2, summary["Maximum"], width,
                       label="Maximum")
        for bars in (bars1, bars2):
            for bar in bars:
                ax.text(bar.get_x() + bar.get_width() / 2,
                        bar.get_height(), f"{bar.get_height():.2f}",
                        ha="center", va="bottom", fontsize=14)
        ax.set_xticks(x)
        ax.set_xticklabels(labels, fontsize=14)
        ax.set_xlabel("Criterion-Value Perturbation", fontsize=16)
        ax.tick_params(axis="y", labelsize=14)
        ax.set_title(title, fontsize=16, fontweight="bold")
        ax.legend(frameon=False, fontsize=14)
    axes[0].set_ylabel("Absolute Rank Change", fontsize=16)
    plt.tight_layout()
    if save_path:
        plt.savefig(save_path, dpi=300, bbox_inches="tight")
    plt.show()

def plot_stochastic_interval_combined(
    ahp_df, pca_df, save_path="AHP_PCA_Stochastic_Interval.png",
    figsize=(12, 5)
):
    fig, axes = plt.subplots(1, 2, figsize=figsize, sharey=True)
    for ax, sim_df, title in zip(
        axes, [ahp_df, pca_df], ["(a) AHP", "(b) PCA"]
    ):
        summary = summarize_stochastic(sim_df)
        x = np.arange(len(summary))
        labels = [f"±{int(v * 100)}%" for v in summary["Perturbation"]]
        cap = 0.08
        ax.vlines(x, summary["P05"], summary["P95"], linewidth=1.5,
                  label="P05–P95")
        ax.vlines(x, summary["Q1"], summary["Q3"], linewidth=8,
                  alpha=0.35, label="IQR")
        ax.scatter(x, summary["Median"], marker="o", s=55,
                   label="Median", zorder=3)
        ax.scatter(x, summary["Mean"], marker="^", s=65,
                   label="Mean", zorder=3)
        for i, row in summary.iterrows():
            ax.hlines(row["P05"], i - cap, i + cap, linewidth=1.5)
            ax.hlines(row["P95"], i - cap, i + cap, linewidth=1.5)
            ax.text(i + 0.05, row["P95"], f"P95={row['P95']:.2f}",
                    fontsize=9, va="bottom")
            ax.text(i + 0.05, row["Q3"], f"Q3={row['Q3']:.2f}",
                    fontsize=9, va="bottom")
            ax.text(i + 0.05, row["Mean"], f"Mean={row['Mean']:.2f}",
                    fontsize=9, va="bottom")
            ax.text(i + 0.05, row["Median"], f"Median={row['Median']:.2f}",
                    fontsize=9, va="top")
            ax.text(i + 0.05, row["Q1"], f"Q1={row['Q1']:.2f}",
                    fontsize=9, va="top")
            ax.text(i + 0.05, row["P05"], f"P05={row['P05']:.2f}",
                    fontsize=9, va="top")
        ax.set_xticks(x)
        ax.set_xticklabels(labels, fontsize=14)
        ax.set_xlabel("Criterion-Value Perturbation", fontsize=16)
        ax.tick_params(axis="y", labelsize=14)
        ax.set_title(title, fontsize=16, fontweight="bold")
        ax.legend(frameon=False, fontsize=12)
    axes[0].set_ylabel("Mean Absolute Rank Change", fontsize=16)
    plt.tight_layout()
    if save_path:
        plt.savefig(save_path, dpi=300, bbox_inches="tight")
    plt.show()

def plot_tornado_oat(
    oat_add_df, oat_sub_df, title="OAT Sensitivity",
    add_cmap=ADD_CMAP, sub_cmap=SUB_CMAP,
    save_path=None,
    legend_fontsize=16,
    legend_title_fontsize=16
):
    df_add, df_sub = oat_add_df.T, oat_sub_df.T
    total = df_add.sum(axis=1) + df_sub.sum(axis=1)
    idx = total.sort_values(ascending=True).index
    add_sorted, sub_sorted = df_add.loc[idx], df_sub.loc[idx]
    def truncate_cmap(name, n, minval=0.2, maxval=0.9):
        cmap = plt.get_cmap(name)
        return mcolors.LinearSegmentedColormap.from_list(
            f"trunc_{name}", cmap(np.linspace(minval, maxval, max(n, 2)))
        )
    add_colors = truncate_cmap(add_cmap, add_sorted.shape[1])
    sub_colors = truncate_cmap(sub_cmap, sub_sorted.shape[1])
    fig, ax = plt.subplots(figsize=(12, 6))
    y = np.arange(len(idx))
    left_tracker = np.zeros(len(idx))
    right_tracker = np.zeros(len(idx))
    for i, delta in enumerate(sub_sorted.columns):
        values = sub_sorted[delta].to_numpy()
        ax.barh(
            y, -values, left=-left_tracker,
            color=sub_colors(i / max(sub_sorted.shape[1] - 1, 1))
        )
        left_tracker += values
    for i, delta in enumerate(add_sorted.columns):
        values = add_sorted[delta].to_numpy()
        ax.barh(
            y, values, left=right_tracker,
            color=add_colors(i / max(add_sorted.shape[1] - 1, 1))
        )
        right_tracker += values
    ax.set_yticks(y)
    plot_labels = []
    plot_labels = [
        CRITERION_NAMES.get(str(c).strip(), str(c).strip())
        for c in idx
    ]
    ax.set_yticklabels(plot_labels, fontsize=16)
    ax.set_xlabel("Total Rank Change", fontsize=20)
    ax.set_title(title, fontsize=15)
    ax.axvline(0, linewidth=0.8)
    ax.invert_yaxis()
    ax.tick_params(axis="x", labelsize=14)
    ax.tick_params(axis="y", length=0)
    for name, spine in ax.spines.items():
        spine.set_visible(name == "bottom")
    add_patches = [
        plt.Rectangle(
            (0, 0), 1, 1,
            color=add_colors(i / max(add_sorted.shape[1] - 1, 1))
        ) for i in range(add_sorted.shape[1])
    ]
    sub_patches = [
        plt.Rectangle(
            (0, 0), 1, 1,
            color=sub_colors(i / max(sub_sorted.shape[1] - 1, 1))
        ) for i in range(sub_sorted.shape[1])
    ]
    labels = (
        [f"+{int(d * 100)}%" for d in add_sorted.columns] +
        [f"-{int(d * 100)}%" for d in sub_sorted.columns]
    )
    ax.legend(
        add_patches + sub_patches, labels,
        title="Weight perturbation", frameon=False,
        fontsize=legend_fontsize, title_fontsize=legend_title_fontsize, loc="upper right"
    )
    plt.tight_layout()
    if save_path:
        plt.savefig(save_path, dpi=300, bbox_inches="tight")
    plt.show()

def plot_lhs_jaccard(lhs_metrics, save_path='LHS_topK_Jaccard.png', title=None, figsize=(5, 5), bar_color='steelblue', edge_color='none', bins=6, label_fontsize=18, tick_fontsize=16, title_fontsize=14, legend_fontsize=16):
    values = lhs_metrics['TopK_Jaccard'].dropna()
    fig, ax = plt.subplots(figsize=figsize)
    ax.hist(values, bins=bins, color=bar_color, edgecolor=edge_color, alpha=0.8)
    median = values.median()
    ax.axvline(median, color='red', linestyle='--', linewidth=2.0, label=f'Median = {median:.3f}')
    ax.set_xlabel(f'Top-{TOP_K} Jaccard Similarity', fontsize=label_fontsize)
    ax.set_ylabel('Number of LHS Realizations', fontsize=label_fontsize)
    if title is None:
        title = f'Distribution of Top-{TOP_K} Ranking Stability'
    ax.set_title(title, fontsize=title_fontsize)
    ax.set_xlim(0, 1)
    ax.tick_params(axis='both', labelsize=tick_fontsize)
    ax.legend(frameon=False, fontsize=legend_fontsize)
    plt.tight_layout()
    if save_path:
        plt.savefig(save_path, dpi=300, bbox_inches='tight')
    plt.show()

def plot_topk_retention(lhs_critical, top_k=TOP_K, save_path='LHS_topK_retention_probability.png', figsize=(8, 7), bar_color='silver', bar_height=0.55, union_fontsize=12, value_fontsize=10, value_fontweight='bold', x_tick_fontsize=16, axis_fontsize=18, title_fontsize=14, x_max=1.1):
    df = lhs_critical.copy()
    df = df.sort_values('Baseline_rank')
    labels = df['Alternative'].astype(str)
    fig, ax = plt.subplots(figsize=figsize)
    bars = ax.barh(labels, df['Probability_remain_in_topK'], height=bar_height, color=bar_color)
    for bar, value in zip(bars, df['Probability_remain_in_topK']):
        ax.text(value + 0.012, bar.get_y() + bar.get_height() / 2, f'{value:.2f}', va='center', ha='left', fontsize=value_fontsize, fontweight=value_fontweight)
    ax.set_xlabel(f'Probability of Remaining in Top {top_k}', fontsize=axis_fontsize)
    ax.set_ylabel('Administrative Units', fontsize=axis_fontsize)
    ax.set_title(f'Stability of Baseline Top-{top_k} Unions', fontsize=title_fontsize)
    ax.set_xlim(0, x_max)
    ax.set_xticks(np.arange(0, 1.01, 0.2))
    ax.tick_params(axis='x', labelsize=x_tick_fontsize)
    ax.tick_params(axis='y', labelsize=union_fontsize, length=0)
    ax.invert_yaxis()
    ax.margins(y=0.01)
    plt.tight_layout()
    if save_path:
        plt.savefig(save_path, dpi=300, bbox_inches='tight')
    plt.show()

def plot_component_topk_stability(component_metrics_df, top_k=TOP_K, save_path='Component_topK_stability.png', figsize=(9, 5), bar_color='slategray', bar_height=0.65, label_fontsize=16, value_fontsize=12, title_fontsize=14):
    df = component_metrics_df.copy()
    df = df.sort_values('TopK_Retained_Count', ascending=True)
    fig, ax = plt.subplots(figsize=figsize)
    bars = ax.barh(df['Scenario'], df['TopK_Retained_Count'], color=bar_color, height=bar_height)
    for bar, value in zip(bars, df['TopK_Retained_Count']):
        ax.text(value + 0.3, bar.get_y() + bar.get_height() / 2, f'{int(value)}/{top_k}', va='center', ha='left', fontsize=value_fontsize, fontweight='bold')
    ax.set_xlabel(f' Top-{top_k} Administrative Units Retained', fontsize=label_fontsize)
    ax.set_ylabel('Structural Scenario of Components', fontsize=title_fontsize)
    ax.set_title(f'Top-{top_k} Stability under Component-Structure Scenarios', fontsize=title_fontsize)
    ax.set_xlim(0, top_k + 4)
    ax.tick_params(axis='both', labelsize=value_fontsize)
    plt.tight_layout()
    if save_path:
        plt.savefig(save_path, dpi=300, bbox_inches='tight')
    plt.show()

def format_excel_workbook(file_path):
    """
    Apply consistent formatting to all result sheets.
    This changes only Excel appearance, not analysis values.
    """
    wb = load_workbook(file_path)
    header_fill = PatternFill('solid', fgColor='D9EAF7')
    header_font = Font(bold=True)
    thin_gray = Side(style='thin', color='D9D9D9')
    for ws in wb.worksheets:
        ws.freeze_panes = 'A2'
        if ws.max_row >= 1 and ws.max_column >= 1:
            ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cell.border = Border(bottom=thin_gray)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical='center', wrap_text=False)
                if isinstance(cell.value, float):
                    cell.number_format = '0.0000'
        for col_idx in range(1, ws.max_column + 1):
            letter = get_column_letter(col_idx)
            max_len = 0
            for cell in ws[letter]:
                value = '' if cell.value is None else str(cell.value)
                max_len = max(max_len, len(value))
            ws.column_dimensions[letter].width = min(max(max_len + 2, 11), 36)
        ws.sheet_view.showGridLines = False
    wb.save(file_path)

if __name__ == '__main__':
    df_full = load_excel_data(FILE_PATH, SHEET_NAME)
    labels = df_full['Union_Name'].astype(str).to_numpy()
    decision_df = df_full.drop(columns=['Alternative_ID', 'Union_Name']).copy()
    criterion_cols = [c for c in decision_df.columns if c != TARGET_COL]
    if len(criterion_cols) != len(CRITERIA):
        raise ValueError(f'Expected {len(CRITERIA)} criteria but found {len(criterion_cols)}: {criterion_cols}')
    ahp_weights = AHP_WEIGHTS / AHP_WEIGHTS.sum()
    pca_w, pca_variance, pca_table = pca_weights(decision_df, TARGET_COL)
    pca_info = pd.DataFrame({'Item': ['Retained components', 'PC1 explained variance', 'PC1 explained variance (%)'], 'Value': [1, pca_variance, pca_variance * 100]})
    comparison_df, baseline_ranks_df, weight_comparison_df = compare_ahp_pca(decision_df, TARGET_COL, ahp_weights, pca_w, CRITERIA, labels)
    component_scores_df, component_metrics_df, component_weights_df, component_topk_changes_df = component_sensitivity(decision_df, TARGET_COL, labels, ahp_weights, CRITERIA)
    ahp_oat_add, ahp_oat_sub, ahp_oat_metrics = oat_weight_sensitivity(decision_df, TARGET_COL, ahp_weights, CRITERIA)
    lhs_metrics, lhs_weights, lhs_critical, lhs_summary, lhs_baseline = lhs_weight_uncertainty(decision_df, TARGET_COL, ahp_weights, CRITERIA, labels, preserve_components=True)
    pca_lhs_metrics, pca_lhs_weights, pca_lhs_critical, pca_lhs_summary, pca_lhs_baseline = lhs_weight_uncertainty(decision_df, TARGET_COL, pca_w, CRITERIA, labels, n_sim=N_LHS, rel_range=LHS_RANGE, seed=LHS_SEED, preserve_components=False)
    stochastic_ahp_df = None
    if RUN_STOCHASTIC_DATA_SA:
        stochastic_ahp_df, stochastic_sim_df = stochastic_data_sensitivity(decision_df, TARGET_COL, ahp_weights, CRITERIA, labels, rel_errors=DATA_ERROR_LEVELS, n_sim=N_DATA_SIM, seed=DATA_SEED, min_values=MIN_VALS, max_values=MAX_VALS)
        stochastic_pca_df, stochastic_pca_sim_df = stochastic_data_sensitivity(decision_df,TARGET_COL,pca_w, CRITERIA,labels,rel_errors=DATA_ERROR_LEVELS,n_sim=N_DATA_SIM,seed=DATA_SEED,min_values=MIN_VALS,max_values=MAX_VALS)
    ahp_stochastic_summary_df = None
    pca_stochastic_summary_df = None
    if RUN_STOCHASTIC_DATA_SA:
        ahp_stochastic_summary_df = summarize_stochastic(stochastic_sim_df)
        pca_stochastic_summary_df = summarize_stochastic(
            stochastic_pca_sim_df
        )
    settings_df = pd.DataFrame({'Setting': ['Weight OAT perturbations', 'LHS simulations', 'LHS relative range', 'LHS seed', 'Highest-risk fraction', 'Top-k', 'Component scenarios', 'Geometric aggregation', 'PCA LHS enabled', 'Unit-specific stochastic data sensitivity'], 'Value': ['±10%, ±20%, ±30%', N_LHS, f'±{LHS_RANGE * 100:.0f}%', LHS_SEED, TOP_PCT, TOP_K, 'H=25%, 33.3%, 50%, 66.7%; S/V asymmetry 30/20 and 20/30', 'S^0.25 × V^0.25 × H^0.50', True, f'{N_DATA_SIM} simulations; ±5%, ±10%, ±15%, ±20%; seed={DATA_SEED}']})
    with pd.ExcelWriter(OUTPUT_FILE, engine='openpyxl') as writer:
        settings_df.to_excel(writer, 'Settings', index=False)
        pca_info.to_excel(writer, 'PCA_Info', index=False)
        pca_table.to_excel(writer, 'PCA_Weights', index=False)
        weight_comparison_df.to_excel(writer, 'AHP_PCA_Weights', index=False)
        comparison_df.to_excel(writer, 'AHP_PCA_Comparison', index=False)
        baseline_ranks_df.to_excel(writer, 'AHP_PCA_Ranks', index=False)
        ahp_oat_add.to_excel(writer, 'AHP_OAT_Add')
        ahp_oat_sub.to_excel(writer, 'AHP_OAT_Sub')
        ahp_oat_metrics.to_excel(writer, 'AHP_OAT_Metrics', index=False)
        lhs_summary.to_excel(writer, 'LHS_Summary', index=False)
        lhs_metrics.to_excel(writer, 'LHS_Metrics', index=False)
        lhs_weights.to_excel(writer, 'LHS_Weights', index=False)
        lhs_critical.to_excel(writer, 'LHS_Critical_Ranks', index=False)
        lhs_baseline.to_excel(writer, 'AHP_Baseline_Ranks', index=False)
        pca_lhs_summary.to_excel(writer, 'PCA_LHS_Summary', index=False)
        pca_lhs_metrics.to_excel(writer, 'PCA_LHS_Metrics', index=False)
        pca_lhs_weights.to_excel(writer, 'PCA_LHS_Weights', index=False)
        pca_lhs_critical.to_excel(writer, 'PCA_LHS_Critical', index=False)
        pca_lhs_baseline.to_excel(writer, 'PCA_Baseline_Ranks', index=False)
        component_scores_df.to_excel(writer, 'Component_Ranks', index=False)
        component_metrics_df.to_excel(writer, 'Component_Metrics', index=False)
        component_weights_df.to_excel(writer, 'Component_Weights', index=False)
        component_topk_changes_df.to_excel(writer, 'Component_TopK_Changes', index=False)
        if RUN_STOCHASTIC_DATA_SA:
            stochastic_ahp_df.to_excel(writer, 'AHP_Stochastic_Data_SA', index=False)
            stochastic_sim_df.to_excel(writer, 'AHP_Stochastic_Simulations', index=False)
            ahp_stochastic_summary_df.to_excel(writer, 'AHP_Stochastic_Summary', index=False)
            stochastic_pca_df.to_excel(writer, 'PCA_Stochastic_Data_SA', index=False)
            stochastic_pca_sim_df.to_excel(writer, 'PCA_Stochastic_Simulations', index=False)
            pca_stochastic_summary_df.to_excel(writer, 'PCA_Stochastic_Summary', index=False)
    plot_tornado_oat(
        ahp_oat_add, ahp_oat_sub,
        title="AHP-SAW Two-Sided Weight Sensitivity",
        save_path=OUTPUT_DIR / "AHP_SAW_weight_OAT.png"
    )
    plot_lhs_jaccard(lhs_metrics, save_path=OUTPUT_DIR / 'AHP_LHS_top36_Jaccard.png', title='AHP-SAW LHS Top-36 Ranking Stability')
    plot_lhs_jaccard(pca_lhs_metrics, save_path=OUTPUT_DIR / 'PCA_LHS_top36_Jaccard.png', title='PCA-SAW LHS Top-36 Ranking Stability')
    plot_topk_retention(lhs_critical, save_path=OUTPUT_DIR / 'AHP_LHS_top36_retention_probability.png')
    plot_topk_retention(pca_lhs_critical, save_path=OUTPUT_DIR / 'PCA_LHS_top36_retention_probability.png')
    plot_component_topk_stability(component_metrics_df, top_k=TOP_K, save_path=OUTPUT_DIR / 'Component_top36_stability.png')
    if RUN_STOCHASTIC_DATA_SA:
        plot_stochastic_data_sensitivity_combined(
            stochastic_ahp_df, stochastic_pca_df,
            save_path=OUTPUT_DIR / "AHP_PCA_Stochastic_Data_Sensitivity.png"
        )
        plot_stochastic_interval_combined(
            stochastic_sim_df, stochastic_pca_sim_df,
            save_path=OUTPUT_DIR / "AHP_PCA_Stochastic_Data_Interval.png"
        )
    format_excel_workbook(OUTPUT_FILE)
    print(f'Results exported to: {OUTPUT_FILE}')
    print(f'PC1 explained variance: {pca_variance:.6f}')
    print(comparison_df.to_string(index=False))
